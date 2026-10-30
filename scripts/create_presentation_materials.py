from pathlib import Path
import json

import matplotlib.pyplot as plt
from matplotlib.ticker import FormatStrFormatter
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import (
    Image,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "paper" / "presentation"
ASSETS = OUT / "assets"
OUT.mkdir(parents=True, exist_ok=True)
ASSETS.mkdir(parents=True, exist_ok=True)


def load(name):
    with (ROOT / "results" / name).open(encoding="utf-8") as handle:
        return json.load(handle)


standardized = load("paired_standardized.json")
raw = load("paired_raw_mv.json")
wide = load(Path("sensitivity_wide") / "summary.json")
classes = ["NORM", "MI", "STTC", "CD", "HYP"]
labels = ["Normal", "MI", "ST/T change", "Conduction\ndisturbance", "Hypertrophy"]


def make_table(data, widths):
    result = Table(data, colWidths=widths, repeatRows=1)
    result.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#123B5D")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#CBD5E1")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1),
         [colors.white, colors.HexColor("#F1F5F9")]),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    return result


def save_charts():
    plt.rcParams.update({"font.family": "DejaVu Sans", "axes.titleweight": "bold"})
    x = range(len(classes))

    fig, ax = plt.subplots(figsize=(9, 4.7))
    ax.bar([i - 0.18 for i in x], [standardized["delta_classwise_ece"][c] for c in classes],
           0.36, label="Standardized", color="#3B82F6")
    ax.bar([i + 0.18 for i in x], [raw["delta_classwise_ece"][c] for c in classes],
           0.36, label="Raw millivolts", color="#EF4444")
    ax.axhline(0, color="#333333", linewidth=0.8)
    ax.set_xticks(list(x), labels)
    ax.set_ylabel("Change in ECE (traced − clean)")
    ax.set_title("Tracing the paper page changes probability calibration")
    ax.yaxis.set_major_formatter(FormatStrFormatter("%.2f"))
    ax.legend(frameon=False)
    ax.grid(axis="y", alpha=0.25)
    fig.tight_layout()
    fig.savefig(ASSETS / "ece_by_class.png", dpi=180)
    plt.close(fig)

    contrast = [raw["delta_classwise_ece"][c] - standardized["delta_classwise_ece"][c] for c in classes]
    fig, ax = plt.subplots(figsize=(9, 4.7))
    colors_bar = ["#94A3B8", "#94A3B8", "#F59E0B", "#DC2626", "#F59E0B"]
    ax.bar(list(x), contrast, color=colors_bar)
    ax.axhline(0, color="#333333", linewidth=0.8)
    ax.set_xticks(list(x), labels)
    ax.set_ylabel("Raw-mV ΔECE − standardized ΔECE")
    ax.set_title("The largest representation-dependent effect is conduction disturbance")
    ax.yaxis.set_major_formatter(FormatStrFormatter("%.2f"))
    ax.grid(axis="y", alpha=0.25)
    for i, value in enumerate(contrast):
        ax.text(i, value + 0.006, f"{value:.3f}", ha="center", fontsize=9, fontweight="bold")
    fig.tight_layout()
    fig.savefig(ASSETS / "contrast.png", dpi=180)
    plt.close(fig)

    fig, ax = plt.subplots(figsize=(9, 4.7))
    values = [0.142, -0.006, -0.005]
    bars = ax.bar(["Before correction", "Logit-only", "Trace-aware"], values,
                  color=["#DC2626", "#10B981", "#059669"])
    ax.axhline(0, color="#333333", linewidth=0.8)
    ax.set_ylabel("CD calibration shift")
    ax.set_title("Validation-fitted recalibration brings the CD shift near zero")
    ax.yaxis.set_major_formatter(FormatStrFormatter("%.3f"))
    ax.grid(axis="y", alpha=0.25)
    for bar, value in zip(bars, values):
        ax.text(bar.get_x() + bar.get_width() / 2, value + (0.008 if value >= 0 else -0.018),
                f"{value:.3f}", ha="center", fontsize=10, fontweight="bold")
    fig.tight_layout()
    fig.savefig(ASSETS / "correction.png", dpi=180)
    plt.close(fig)


def make_pdf():
    path = OUT / "ecg_digitization_research_guide.pdf"
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="TitleCenter", parent=styles["Title"], alignment=TA_CENTER,
                              textColor=colors.HexColor("#123B5D"), fontSize=24, leading=29))
    styles.add(ParagraphStyle(name="Subtitle", parent=styles["Normal"], alignment=TA_CENTER,
                              fontSize=12, textColor=colors.HexColor("#475569"), leading=17))
    styles.add(ParagraphStyle(name="H", parent=styles["Heading2"], textColor=colors.HexColor("#123B5D"),
                              spaceBefore=12, spaceAfter=6))
    styles.add(ParagraphStyle(name="Body2", parent=styles["BodyText"], fontSize=10.2, leading=14))
    styles.add(ParagraphStyle(name="Small", parent=styles["BodyText"], fontSize=8.5, leading=11,
                              textColor=colors.HexColor("#475569")))
    story = [
        Paragraph("ECG Paper Digitization and Probability Reliability", styles["TitleCenter"]),
        Spacer(1, 0.15 * inch),
        Paragraph("A beginner-friendly guide to the problem, experiment, findings, and novelty",
                  styles["Subtitle"]),
        Spacer(1, 0.25 * inch),
        Paragraph("<b>One-sentence takeaway:</b> When a clean digital ECG is printed, creased, and digitized again, the classifier's probabilities can become miscalibrated. The effect is diagnosis-specific and is much larger when the model uses raw millivolt amplitudes.", styles["Body2"]),
        Spacer(1, 0.2 * inch),
        Image(str(ASSETS / "ece_by_class.png"), width=6.9 * inch, height=3.6 * inch),
        Paragraph("Figure 1. Positive values mean tracing increased calibration error.", styles["Small"]),
        PageBreak(),
        Paragraph("1. The problem in plain language", styles["H"]),
        Paragraph("An ECG can exist as a clean digital waveform or as a paper printout. In many real workflows, an old paper ECG is scanned or photographed and then converted back into a digital signal. That conversion introduces wrinkles, creases, missing ink, scale changes, and other artifacts.", styles["Body2"]),
        Paragraph("A classifier may still rank diseases reasonably well after this conversion, but its probabilities can become unreliable. A prediction of 90% should mean that roughly 90% of similar predictions are correct. Calibration error measures how far the reported probabilities are from that reality.", styles["Body2"]),
        Paragraph("<b>Research question:</b> Does tracing a synthetic PTB-XL paper page change classwise calibration, and does the change differ by diagnosis?", styles["Body2"]),
        Paragraph("2. What the repository contains", styles["H"]),
        Paragraph("The repository is an end-to-end, reproducible experiment: data validation, label creation, GPU model training, paper-page rendering, waveform tracing, paired evaluation, bootstrap uncertainty, mechanism checks, correction experiments, and a manuscript draft.", styles["Body2"]),
        Paragraph("The classifier is a compact residual 1D network. It is not presented as a new architecture; it is the controlled measuring instrument used to compare clean and traced ECGs.", styles["Body2"]),
        Paragraph("3. Data and experimental design", styles["H"]),
        Paragraph("We use PTB-XL 1.0.3: 12 leads, 100 Hz, 10-second recordings, and five diagnostic superclasses: NORM, MI, STTC, CD, and HYP. Folds 1–8 train, fold 9 selects the checkpoint, and fold 10 is evaluated once. Patient-level bootstrap resampling uses 1,000 resamples.", styles["Body2"]),
        Paragraph("Two identical training arms differ only in input representation: per-lead standardization versus absolute raw millivolts. Both use seed 42 and the same architecture and splits.", styles["Body2"]),
        PageBreak(),
        Paragraph("4. The paper-to-digital experiment", styles["H"]),
        Paragraph("Each test ECG is rendered as a synthetic paper page with red grid, wrinkles, and six creases. The page has no header, calibration pulse, lead names, or handwriting. A column-based dark-ink max-finder recovers the waveform. A training-record correlation gate must pass before test scoring.", styles["Body2"]),
        Paragraph("The primary outcome is ΔECE = ECE(traced waveform) − ECE(clean waveform). Positive ΔECE means the paper/digitization path made the probabilities less reliable.", styles["Body2"]),
        Image(str(ASSETS / "contrast.png"), width=6.9 * inch, height=3.6 * inch),
        Paragraph("Figure 2. Difference between the two input representations.", styles["Small"]),
        Paragraph("5. Main result", styles["H"]),
        Paragraph("The raw-millivolt model showed a conduction-disturbance calibration shift of +0.142. The standardized model changed by −0.019. Their contrast was approximately +0.161, with a patient-bootstrap interval of about 0.136 to 0.171. The same direction and magnitude were reproduced with a wider network (approximately +0.154).", styles["Body2"]),
        PageBreak(),
        Paragraph("6. Did we improve anything?", styles["H"]),
        Paragraph("The experiment did not improve clean classification accuracy. Instead, it revealed a deployment failure and demonstrated a practical fix. A recalibration map fitted on validation data reduced the raw-millivolt conduction-disturbance shift from +0.142 to approximately −0.006 using logits alone, and to approximately −0.005 with trace features.", styles["Body2"]),
        Image(str(ASSETS / "correction.png"), width=6.9 * inch, height=3.6 * inch),
        Paragraph("Figure 3. Correction is applied after the classifier and does not retrain it.", styles["Small"]),
        Paragraph("7. Why the result is meaningful", styles["H"]),
        Paragraph("The effect is not explained only by poor signal reconstruction: CD records had similar SNR to records without CD. Oracle amplitude matching removed only part of the shift. The result therefore indicates that the interaction between trace artifacts and the model's learned voltage representation matters.", styles["Body2"]),
        Paragraph("8. Novelty and limitations", styles["H"]),
        Paragraph("The novelty is the paired, diagnosis-specific study of probability calibration after ECG paper digitization, together with the comparison of standardized versus raw-millivolt inputs and validation-only recalibration. The claim is limited to synthetic PTB-XL creased pages; real photographs and clinical deployment remain future work.", styles["Body2"]),
        PageBreak(),
        Paragraph("9. Ten-minute presentation script", styles["H"]),
        Paragraph("<b>Minute 0–1:</b> Introduce ECGs, paper records, and why probabilities—not only labels—matter.", styles["Body2"]),
        Paragraph("<b>Minute 1–2:</b> State the question: does printing and tracing change classwise calibration?", styles["Body2"]),
        Paragraph("<b>Minute 2–3:</b> Explain PTB-XL, the five classes, and the train/validation/test split.", styles["Body2"]),
        Paragraph("<b>Minute 3–4:</b> Explain the two arms: standardized versus raw millivolt inputs.", styles["Body2"]),
        Paragraph("<b>Minute 4–5:</b> Show the clean-to-paper-to-traced pipeline and the trace-quality gate.", styles["Body2"]),
        Paragraph("<b>Minute 5–7:</b> Present the main graph: raw millivolt CD ΔECE is +0.142, versus −0.019 standardized; contrast +0.161.", styles["Body2"]),
        Paragraph("<b>Minute 7–8:</b> Explain that AUROC and SNR alone do not capture this probability problem.", styles["Body2"]),
        Paragraph("<b>Minute 8–9:</b> Show recalibration reducing the CD shift near zero.", styles["Body2"]),
        Paragraph("<b>Minute 9–10:</b> State the contribution and limitations: calibration must be validated on the same digitization pipeline; current claim is synthetic pages only.", styles["Body2"]),
    ]
    doc = SimpleDocTemplate(str(path), pagesize=A4, rightMargin=0.6 * inch, leftMargin=0.6 * inch,
                            topMargin=0.55 * inch, bottomMargin=0.55 * inch)
    doc.build(story)
    return path


def make_detailed_pdf():
    path = OUT / "ecg_digitization_research_guide.pdf"
    styles = getSampleStyleSheet()
    paragraph = lambda text, style: Paragraph(text, style)
    table = make_table
    styles.add(ParagraphStyle(name="DTitle", parent=styles["Title"], alignment=TA_CENTER,
                              textColor=colors.HexColor("#123B5D"), fontSize=25, leading=30))
    styles.add(ParagraphStyle(name="DSub", parent=styles["Normal"], alignment=TA_CENTER,
                              fontSize=12, textColor=colors.HexColor("#475569"), leading=17))
    styles.add(ParagraphStyle(name="DH", parent=styles["Heading2"], textColor=colors.HexColor("#123B5D"),
                              fontSize=16, leading=20, spaceBefore=13, spaceAfter=7))
    styles.add(ParagraphStyle(name="DH3", parent=styles["Heading3"], textColor=colors.HexColor("#0F766E"),
                              fontSize=11.5, leading=14, spaceBefore=8, spaceAfter=4))
    styles.add(ParagraphStyle(name="DBody", parent=styles["BodyText"], fontSize=9.6, leading=13.3,
                              spaceAfter=6))
    styles.add(ParagraphStyle(name="DNote", parent=styles["BodyText"], fontSize=8.3, leading=10.5,
                              textColor=colors.HexColor("#475569"), spaceAfter=4))
    styles.add(ParagraphStyle(name="DCallout", parent=styles["BodyText"], fontSize=11, leading=15,
                              backColor=colors.HexColor("#E0F2FE"), borderColor=colors.HexColor("#38BDF8"),
                              borderWidth=0.7, borderPadding=8, spaceBefore=5, spaceAfter=8))

    def heading(text):
        return paragraph(text, styles["DH"])

    def subheading(text):
        return paragraph(text, styles["DH3"])

    def body(text):
        return paragraph(text, styles["DBody"])

    def note(text):
        return paragraph(text, styles["DNote"])

    def callout(text):
        return paragraph(text, styles["DCallout"])

    story = [
        paragraph("ECG Digitization Research Guide", styles["DTitle"]),
        Spacer(1, 0.1 * inch),
        paragraph("From absolute beginner to a confident 30-minute presentation", styles["DSub"]),
        Spacer(1, 0.18 * inch),
        callout("<b>Core message:</b> Printing and tracing an ECG can leave classification ranking partly usable while making the model's probabilities less trustworthy. The strongest effect was a +0.142 calibration shift for conduction disturbance in the raw-millivolt model. A validation-fitted recalibration reduced that shift to approximately zero."),
        Image(str(ASSETS / "ece_by_class.png"), width=6.9 * inch, height=3.6 * inch),
        note("Figure 1. Positive bars mean the traced page has higher calibration error than the clean waveform. Blue: standardized model. Red: raw millivolt model."),
        body("<b>How to use this guide:</b> Read Sections 1–8 before presenting. Use the speaker notes in Section 14 to pace a 30-minute talk. The final sections provide exact evidence locations, limitations, and answers to likely questions. All numbers in this guide are taken from the saved repository results."),
        PageBreak(),

        heading("1. The 30-second explanation"),
        body("An ECG is a recording of the heart's electrical activity. Modern ECG systems store the signal digitally, but many records exist on paper. If a paper ECG is scanned, photographed, and converted back to a waveform, the recovered signal is not identical to the original. It can contain distortions from creases, wrinkles, grid lines, missing ink, scale errors, and imperfect tracing."),
        body("A machine-learning classifier receives the recovered waveform and predicts diagnostic categories. It outputs probabilities such as 0.80 for conduction disturbance. The important question is not only whether the model ranks positive cases above negative cases. We must also ask whether 0.80 really corresponds to an approximately 80% observed frequency. That property is called calibration."),
        body("This repository measures the change in calibration caused by the paper-to-digital pipeline, separately for five diagnostic classes and for two voltage representations."),
        callout("<b>One-sentence presentation opening:</b> We studied whether digitizing a printed ECG changes not only what an AI model predicts, but also how much we should trust the model's probability."),

        heading("2. ECG basics for a beginner"),
        subheading("What is an ECG?"),
        body("An electrocardiogram records voltage differences generated by the heart over time. A standard ECG uses 12 leads. Each lead is a different view of the same cardiac electrical activity. The result is a time series: voltage on the vertical axis and time on the horizontal axis."),
        subheading("What is a waveform?"),
        body("The digital ECG is an array of numbers. In this project each recording contains 12 leads sampled at 100 Hz for 1,000 samples, approximately 10 seconds. A computer model does not see a diagnosis or a paper drawing; it sees these numerical sequences."),
        subheading("What is a paper ECG?"),
        body("A paper ECG is a visual representation of the waveform. The grid provides scale, while the dark trace carries the signal. Digitization attempts to reverse this process: find the trace in the image and estimate the original voltage at each time point."),
        subheading("Why digitization can fail"),
        body("The printed page may be wrinkled, creased, rotated, cropped, faded, noisy, photographed under uneven lighting, or scanned at an unknown scale. A digitizer can reconstruct a waveform that looks plausible but differs in amplitude, baseline, timing, or morphology."),

        heading("3. The problem the research solves"),
        body("Earlier evaluation of an ECG classifier often starts with a clean digital waveform and ends with a class label or AUROC. That is incomplete for a probability-based system. A model can keep reasonable AUROC while its probabilities become overconfident or underconfident after digitization."),
        body("This project treats the paper-to-digital process as a deployment transformation. We compare the same frozen model on the same patient's clean waveform and traced waveform. The paired comparison isolates the effect of tracing more directly than comparing unrelated records."),
        body("<b>Primary research question:</b> Does tracing a synthetic PTB-XL paper page change classwise calibration, and does that change differ by diagnosis?"),
        body("<b>What is not the question:</b> We are not claiming a new neural-network architecture, not measuring clinical outcomes, and not claiming that synthetic pages represent every real photograph or historical thermal-paper record."),

        heading("4. The repository at a glance"),
        table([
            ["Repository area", "Role in the research"],
            ["src/", "Data validation, training, rendering, digitization, evaluation, metrics, corrections, and sensitivity analyses."],
            ["configs/", "Locked dataset, model, preprocessing, split, and paired-scoring settings."],
            ["results/", "Saved baselines, paired scores, bootstrap intervals, correction outputs, SNR analysis, and sensitivity results."],
            ["tests/", "Checks for data, labels, leakage, training, transfer, and trace behavior."],
            ["paper/", "Manuscript, draft, figures, and presentation materials."],
            ["README.md / PROTOCOL.md", "Reproduction instructions and the locked scientific question."],
        ], [1.35 * inch, 5.95 * inch]),
        body("The repository is deliberately organized so a future AI or researcher can distinguish source code, configuration, primary results, secondary results, and blocked work."),
        PageBreak(),

        heading("5. Dataset and diagnostic labels"),
        body("The experiment uses PTB-XL 1.0.3. It is not included in the repository because the dataset has its own license and distribution terms; the repository contains the code and configuration needed to validate and use it."),
        table([
            ["Code", "Meaning", "Plain-language interpretation"],
            ["NORM", "Normal", "No target abnormality in the mapped diagnostic statements."],
            ["MI", "Myocardial infarction", "Evidence associated with heart attack."],
            ["STTC", "ST/T change", "Changes involving ST segments or T waves."],
            ["CD", "Conduction disturbance", "Abnormal electrical conduction through the heart."],
            ["HYP", "Hypertrophy", "Thickening-related electrical patterns."],
        ], [0.8 * inch, 1.8 * inch, 4.7 * inch]),
        body("The labels are multi-label: one recording can belong to more than one superclass. Records without diagnostic targets are not silently treated as normal. They remain visible in the Stage 1 data manifest but are excluded from supervised Stage 2 training and evaluation."),
        table([
            ["Split", "Purpose", "Labeled records", "Unlabeled removed from supervised stage"],
            ["Folds 1–8", "Train model parameters", "17,084", "334"],
            ["Fold 9", "Select checkpoint and fit corrections", "2,146", "37"],
            ["Fold 10", "Final paired evaluation", "2,158", "40"],
        ], [1.0 * inch, 2.45 * inch, 1.4 * inch, 2.45 * inch]),
        body("Patient separation is checked across the three splits. The test set is not used to select the model or tune the correction."),

        heading("6. The complete pipeline"),
        body("<b>Stage A — Data preparation.</b> Validate checksums, waveform dimensions, lead order, sampling rate, units, finite values, labels, and patient-separated folds."),
        body("<b>Stage B — Clean model training.</b> Train two otherwise identical classifiers. One receives per-lead standardized data; the other receives raw millivolts. Select each checkpoint using fold-9 macro-AUROC."),
        body("<b>Stage C — Paper rendering.</b> Convert the clean waveform into a synthetic paper page using the locked ECG-Image-Kit renderer. The primary condition includes red grid, wrinkles, and six 45-degree creases."),
        body("<b>Stage D — Trace recovery.</b> Locate lead boxes and recover the dark waveform column by column. The recovered signal is an estimated waveform, not a copy of the plotted pixels."),
        body("<b>Stage E — Quality gate.</b> Trace a training record first. Every lead must correlate at least 0.80 with its clean version. Only after this gate may fold 10 be scored."),
        body("<b>Stage F — Paired scoring.</b> Run both frozen models on both clean and traced versions of the same fold-10 records. Compute ECE, AUROC, entropy, Brier score, and reconstruction SNR."),
        body("<b>Stage G — Uncertainty and robustness.</b> Use patient-cluster bootstrap, compare trace SNR by class, perform oracle scale matching, fit validation-only correction maps, and repeat the primary contrast with a wider network."),
        body("<b>Stage H — Reporting.</b> Keep the five classwise ΔECE values primary. Treat corrections, SNR mechanism analysis, tracer comparison, and width sensitivity as secondary evidence."),
        PageBreak(),

        heading("7. Why there are two model arms"),
        subheading("Standardized arm"),
        body("Each lead is z-scored over time. This removes absolute voltage scale and emphasizes the shape and relative variation of the waveform."),
        subheading("Raw-millivolt arm"),
        body("The original voltage scale is preserved. This can retain useful physical information, but it can also make the model sensitive to amplitude changes introduced by a paper renderer or digitizer."),
        body("Both arms use the same seed, architecture, split, loss, optimizer, schedule, and scoring procedure. This is a controlled comparison: the main intentional difference is voltage representation."),
        table([
            ["Property", "Standardized", "Raw millivolts"],
            ["Clean macro-AUROC", "0.9076", "0.9259"],
            ["Clean ECE", "0.0398", "0.0309"],
            ["Clean Brier", "0.0913", "0.0839"],
            ["Meaning", "Slightly weaker clean discrimination here, but less CD calibration shift after tracing.", "Stronger clean baseline here, but much larger CD calibration shift after tracing."],
        ], [2.0 * inch, 2.6 * inch, 2.7 * inch]),
        callout("<b>Important interpretation:</b> Better clean AUROC or clean ECE does not guarantee better deployment calibration after a paper/digitization transformation."),

        heading("8. Metrics explained from zero"),
        subheading("AUROC: ranking quality"),
        body("AUROC asks whether positive cases tend to receive higher scores than negative cases. It is a ranking metric. An AUROC near 1.0 is strong; 0.5 is approximately random ranking. AUROC does not require the numerical probabilities to be trustworthy."),
        subheading("Calibration and ECE: probability quality"),
        body("Calibration asks whether predicted probabilities match observed frequencies. ECE groups probabilities into 15 equal-width bins and compares the average predicted probability with the observed positive fraction in each bin. Lower ECE is better. The project reports the change, ΔECE, rather than only an absolute number."),
        subheading("Brier score: probability error"),
        body("Brier score averages squared probability error over record/class pairs. Lower is better. It combines calibration and discrimination effects, so the classwise ΔECE remains the primary measure for this research question."),
        subheading("Entropy: uncertainty behavior"),
        body("Binary entropy measures uncertainty in a probability. It is highest near 0.5 and low near 0 or 1. A change in entropy indicates that the tracing pipeline changes how uncertain or decisive the model's probabilities are."),
        subheading("SNR: waveform reconstruction quality"),
        body("Signal-to-noise ratio compares the recovered waveform with the clean waveform under the PhysioNet Challenge 2024 alignment rules. Higher is better, but SNR alone does not tell us whether classifier probabilities are calibrated."),
        subheading("Bootstrap interval"),
        body("The patient-cluster bootstrap repeatedly resamples patients and recomputes the effect. The reported 2.5th–97.5th percentile interval describes the observed uncertainty under this resampling procedure. It is not a guarantee about every future dataset."),
        PageBreak(),

        heading("9. Primary results"),
        body("The primary test contains 2,158 labeled fold-10 records. The average reconstruction SNR is 9.06 dB in both primary arms because the page and digitization pipeline are shared."),
        Image(str(ASSETS / "ece_by_class.png"), width=6.9 * inch, height=3.6 * inch),
        note("Figure 2. Classwise change in ECE after tracing. A positive value means calibration worsened."),
        table([
            ["Class", "Standardized ΔECE", "Raw-mV ΔECE", "Interpretation"],
            ["NORM", "+0.056", "+0.066", "Calibration worsened in both arms."],
            ["MI", "+0.075", "+0.077", "Similar increase in both arms."],
            ["STTC", "+0.005", "+0.041", "Larger increase in raw millivolts."],
            ["CD", "−0.019", "<b>+0.142</b>", "<b>Largest representation-dependent effect.</b>"],
            ["HYP", "+0.034", "+0.068", "Increase in both arms, larger in raw millivolts."],
        ], [0.9 * inch, 1.45 * inch, 1.45 * inch, 3.5 * inch]),
        body("The standardized CD point estimate is slightly negative, while the raw-millivolt CD estimate is strongly positive. This is why the paper emphasizes the raw-minus-standardized contrast as the robust finding rather than claiming that every preprocessing arm behaves identically."),

        heading("10. The central result: raw minus standardized"),
        Image(str(ASSETS / "contrast.png"), width=6.9 * inch, height=3.6 * inch),
        note("Figure 3. The difference between the two preprocessing arms on the same traced-page experiment."),
        body("For conduction disturbance, the raw-millivolt calibration shift was approximately +0.142 and the standardized shift was approximately −0.019. Their difference was approximately <b>+0.161</b>. The patient-bootstrap interval was approximately <b>0.136 to 0.171</b>, excluding zero."),
        body("The other raw-minus-standardized contrasts were approximately +0.010 for NORM, +0.002 for MI, +0.037 for STTC, and +0.034 for HYP. The effect is therefore diagnosis-specific, with CD clearly dominant."),
        callout("<b>Presentation sentence:</b> The same page trace was much more damaging to CD probability calibration when the model had learned from absolute voltage than when it had learned from standardized leads."),
        PageBreak(),

        heading("11. What happened to AUROC and uncertainty?"),
        body("The tracing pipeline also reduced AUROC. In the standardized arm, classwise AUROC changes were approximately −0.033, −0.041, −0.042, −0.038, and −0.078 for NORM, MI, STTC, CD, and HYP. In the raw-millivolt arm they were approximately −0.050, −0.044, −0.048, −0.039, and −0.082."),
        body("This matters because AUROC and calibration answer different questions. AUROC asks whether cases are ranked correctly. Calibration asks whether the numerical probability is reliable. The project shows why reporting AUROC alone can miss an important deployment failure."),
        table([
            ["Metric", "Question answered", "What this project found"],
            ["AUROC", "Are positives ranked above negatives?", "Usually decreased after tracing."],
            ["ECE", "Do probabilities match observed frequencies?", "Large diagnosis-specific raw-mV CD increase."],
            ["Entropy", "Did the model's uncertainty change?", "Raw-mV CD entropy increased by about 0.187 nats."],
            ["SNR", "How similar is recovered waveform to clean waveform?", "Mean test SNR about 9.06 dB."],
        ], [1.1 * inch, 3.0 * inch, 3.2 * inch]),
        body("The raw-mV CD entropy change was approximately +0.187 nats. This indicates the paper trace changed the model's uncertainty distribution, not just the final binary ranking."),

        heading("12. Mechanism investigation"),
        subheading("Is CD simply reconstructed much worse?"),
        body("Not enough to explain the main effect. Mean SNR was approximately 8.83 dB for records with CD and 9.13 dB without CD. The difference is modest compared with the size of the calibration contrast. Some other classes, especially HYP, showed larger SNR differences without producing the same raw-versus-standardized effect."),
        subheading("Is the problem only amplitude scale?"),
        body("An oracle scale check affine-matched every traced lead to that record's clean mean and standard deviation. The raw-millivolt CD shift remained approximately +0.125, with an interval around 0.106–0.136. Scale mismatch contributes, but it is not the whole mechanism."),
        subheading("Does a wider network reproduce the result?"),
        body("Yes. A pre-specified width-64 sensitivity produced a raw-minus-standardized CD contrast of approximately +0.154, with an interval around 0.136–0.169. The agreement rule was satisfied: same sign and interval excluding zero."),
        Image(str(ASSETS / "correction.png"), width=6.9 * inch, height=3.6 * inch),
        note("Figure 4. The correction experiment is secondary: it evaluates whether validation-only recalibration can reduce the observed raw-mV CD shift."),
        PageBreak(),

        heading("13. Correction experiment: what actually improved?"),
        body("The classifier was not retrained. Instead, a correction map was fitted on validation fold 9 and applied once to test fold 10. This respects the test-set lock."),
        subheading("Logit-only correction"),
        body("A separate logistic map was fitted for each class using the frozen model logit. For the raw-millivolt arm, the corrected shifts were approximately −0.009, +0.005, +0.005, −0.006, and +0.003 for NORM, MI, STTC, CD, and HYP."),
        subheading("Trace-aware correction"),
        body("A second map used the frozen logit plus traced-lead standard deviation and fraction of columns with no ink. For CD it moved the shift to approximately −0.005, with an interval approximately −0.018 to +0.012. Traced-page macro-AUROC stayed within the pre-specified 0.01 tolerance."),
        body("The main practical improvement is therefore <b>probability reliability</b>, not clean accuracy. A deployment pipeline can preserve the classifier and add calibration appropriate to the digitized input domain."),
        callout("<b>Correct wording:</b> We found and corrected a digitization-induced calibration shift. We did not prove that the correction is clinically validated or universally transferable."),

        heading("14. Novelty and contribution"),
        body("The classifier architecture is not claimed as novel. The page tracer is also an existing tool. The contribution is the controlled experiment and its interpretation."),
        table([
            ["Contribution", "Why it matters"],
            ["Paired clean-versus-traced evaluation", "The same records isolate the effect of the paper pipeline."],
            ["Classwise calibration as primary outcome", "Finds probability failures that AUROC-only studies can miss."],
            ["Standardized versus raw-mV arms", "Shows voltage representation changes vulnerability."],
            ["Mechanism checks", "Separates reconstruction quality, amplitude scale, and model sensitivity."],
            ["Validation-only correction", "Demonstrates a practical route to reduce the shift."],
            ["Locked protocol and bootstrap", "Reduces test-set tuning and makes the evidence auditable."],
        ], [2.3 * inch, 5.0 * inch]),
        PageBreak(),

        heading("15. What the project does not claim"),
        body("The current evidence is limited to synthetic PTB-XL paper pages with the locked creased condition. It is not a clinical trial and does not establish safety, diagnostic benefit, or deployment readiness."),
        table([
            ["Supported", "Not supported yet"],
            ["A synthetic creased trace changed classwise calibration.", "Generalization to every scanner, camera, or digitizer."],
            ["Raw millivolts had a larger CD calibration shift than standardized input.", "That raw millivolts are always worse."],
            ["Validation-only recalibration reduced the observed shift.", "Clinical validity of the correction."],
            ["Width-64 sensitivity reproduced the main contrast.", "Universal reproduction across seeds and architectures."],
            ["Primary test-page SNR was about 9.06 dB.", "Real photographs, thermal paper, or phone images."],
        ], [3.65 * inch, 3.65 * inch]),
        subheading("Blocked experiment"),
        body("The external color-scan test was not scored because the required color-scan variant identifier was unset. The gate is recorded in results/real_scan/gate.json. This must be described as blocked, not as a negative result."),

        heading("16. Thirty-minute presentation plan"),
        table([
            ["Time", "Topic", "Message to deliver"],
            ["0:00–2:00", "Motivation", "Paper ECGs must be digitized, and probability trust matters beyond labels."],
            ["2:00–4:00", "ECG basics", "A 12-lead ECG is a multichannel voltage time series; paper is a visual encoding."],
            ["4:00–6:00", "Problem", "Tracing introduces artifacts that can change model probabilities."],
            ["6:00–8:00", "Question", "Does classwise calibration change, and does diagnosis matter?"],
            ["8:00–11:00", "Dataset", "PTB-XL, five classes, patient-separated folds, multi-label targets."],
            ["11:00–14:00", "Two arms", "Same model and split; standardized leads versus raw millivolts."],
            ["14:00–17:00", "Pipeline", "Clean waveform → synthetic creased page → trace recovery → paired scoring."],
            ["17:00–20:00", "Metrics", "AUROC is ranking; ECE is probability reliability; SNR is reconstruction."],
            ["20:00–23:00", "Main result", "Raw-mV CD +0.142, standardized CD −0.019, contrast +0.161."],
            ["23:00–25:00", "Mechanism", "Not explained only by SNR or amplitude scale; wider model agrees."],
            ["25:00–27:00", "Correction", "Validation recalibration moves CD shift near zero."],
            ["27:00–29:00", "Novelty", "Diagnosis-specific calibration study, representation comparison, practical correction."],
            ["29:00–30:00", "Limits / close", "Synthetic scope only; calibration must be checked on the actual input pipeline."],
        ], [1.0 * inch, 1.5 * inch, 4.8 * inch]),
        PageBreak(),

        heading("17. Speaker notes: simple explanations"),
        subheading("Explain calibration with a weather analogy"),
        body("If a weather app says there is a 70% chance of rain on ten similar days, it should rain on roughly seven of them. A model can be good at ranking rainy days above dry days but still be badly calibrated if its 70% predictions rain only on five days."),
        subheading("Explain the key result without jargon"),
        body("When the model was trained on raw voltage, the traced paper signal made its CD probabilities substantially less trustworthy. When the model was trained on standardized signals, the same paper trace caused much less of this CD-specific problem."),
        subheading("Explain why this is positive research"),
        body("We did not improve clean AUROC. We improved the scientific understanding of a hidden deployment failure and showed that a validation-only correction can reduce the probability error. Finding and fixing a failure mode is a meaningful improvement even when the classifier itself is unchanged."),
        subheading("Explain the primary number"),
        body("A CD contrast of +0.161 means the raw-millivolt arm's tracing-induced calibration shift was 0.161 larger than the standardized arm's shift on the ECE scale. The bootstrap interval does not include zero, so the observed contrast is not consistent with no difference under the specified resampling analysis."),

        heading("18. Likely questions and answers"),
        table([
            ["Question", "Answer"],
            ["Did accuracy improve?", "No. Clean classification was not the main improvement. The contribution is identifying and reducing probability calibration error after tracing."],
            ["What is the strongest result?", "The raw-minus-standardized CD ΔECE contrast: about +0.161, interval about 0.136–0.171."],
            ["Why compare two input types?", "To test whether preserving absolute voltage makes the model more sensitive to paper-tracing scale and morphology changes."],
            ["Could this only be bad image quality?", "Not fully. CD SNR was close to non-CD SNR, and oracle scale matching left most of the shift."],
            ["Was the test set used to tune the correction?", "No. The correction was fitted on fold 9 and applied once to fold 10."],
            ["Can we claim real-world clinical impact?", "No. The current claim is limited to synthetic PTB-XL creased pages."],
            ["Why not report only AUROC?", "AUROC measures ranking, not whether probabilities correspond to observed frequencies."],
            ["What is future work?", "Complete the external color-scan gate, add seeds/independent data, and validate correction transfer."],
        ], [2.35 * inch, 4.95 * inch]),
        PageBreak(),

        heading("19. Evidence map for a future AI"),
        table([
            ["Need to verify", "Read this file"],
            ["Research question and estimands", "PROTOCOL.md"],
            ["Reproduction commands and scope", "README.md"],
            ["Historical decisions and status", "CHANGELOG.md"],
            ["Standardized paired primary result", "results/paired_standardized.json"],
            ["Raw-mV paired primary result", "results/paired_raw_mv.json"],
            ["Clean baseline metrics", "results/stage2/REPORT.md and results/stage2_raw/REPORT.md"],
            ["Correction details", "results/corrector/raw_mv.json and results/corrector/logit_only.json"],
            ["SNR mechanism analysis", "results/mechanism/snr_by_class.json"],
            ["Width sensitivity", "results/sensitivity_wide/summary.json"],
            ["External scan status", "results/real_scan/gate.json"],
            ["Implementation", "src/train.py, src/eval.py, src/paired.py, src/digitize.py, src/metrics.py"],
        ], [2.65 * inch, 4.65 * inch]),
        body("When extending the project, preserve the locked primary outcome, creased-page condition, fold assignment, patient bootstrap, one-score-per-arm rule, and distinction between primary and secondary outcomes. Do not tune against fold 10 or overwrite locked results just because a rerun is numerically different."),

        heading("20. Final takeaway"),
        callout("<b>Research conclusion:</b> ECG paper digitization is not only a waveform reconstruction problem. It can change the reliability of diagnostic probabilities in a diagnosis-specific and preprocessing-dependent way."),
        body("The project found a large conduction-disturbance calibration shift for the raw-millivolt model (+0.142 after tracing), a much smaller standardized-arm shift (−0.019), and a raw-minus-standardized contrast of approximately +0.161 with a bootstrap interval excluding zero. The contrast was reproduced with a wider network. Signal quality and amplitude scale explained only part of the effect. A validation-only recalibration reduced the CD shift to approximately −0.005 to −0.006."),
        body("<b>Final practical recommendation:</b> Any AI system intended to read digitized ECGs should be evaluated and calibrated on digitized ECGs produced by the same kind of input pipeline. Clean-waveform AUROC alone is not enough."),
        note("Generated from the repository's verified protocol and saved result artifacts. Synthetic-page scope and blocked external-scan work must remain explicit in any presentation or manuscript."),
    ]

    doc = SimpleDocTemplate(
        str(path), pagesize=A4, rightMargin=0.55 * inch, leftMargin=0.55 * inch,
        topMargin=0.48 * inch, bottomMargin=0.48 * inch,
        title="ECG Digitization Research Guide",
        author="Paper ECG Uncertainty project",
    )
    doc.build(story)
    return path


def add_text(slide, text, left, top, width, height, size=22, color=(30, 41, 59), bold=False):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.06)
    tf.margin_right = Inches(0.06)
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = RGBColor(*color)
    return box


def make_ppt():
    path = OUT / "ecg_digitization_research_presentation.pptx"
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    def slide(title, subtitle=None):
        s = prs.slides.add_slide(blank)
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = RGBColor(248, 250, 252)
        add_text(s, title, 0.55, 0.35, 12.2, 0.55, size=28, color=(18, 59, 93), bold=True)
        if subtitle:
            add_text(s, subtitle, 0.58, 0.95, 12, 0.4, size=14, color=(71, 85, 105))
        return s

    s = prs.slides.add_slide(blank)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = RGBColor(18, 59, 93)
    add_text(s, "ECG Paper Digitization", 0.7, 1.35, 12, 0.75, size=38, color=(255, 255, 255), bold=True)
    add_text(s, "How printing and tracing can change probability reliability", 0.75, 2.25, 11.5, 0.6, size=23, color=(219, 234, 254))
    add_text(s, "10-minute research presentation | PTB-XL | synthetic creased pages", 0.78, 5.85, 11, 0.4, size=15, color=(226, 232, 240))

    s = slide("1. Why this problem matters", "A model can be accurate at ranking cases but wrong about how confident it should be.")
    add_text(s, "Clean digital ECG", 0.8, 2.0, 3.0, 0.4, size=20, bold=True)
    add_text(s, "→", 4.0, 2.0, 0.6, 0.4, size=28, color=(14, 116, 144), bold=True)
    add_text(s, "Paper / scan / trace", 4.75, 2.0, 3.2, 0.4, size=20, bold=True)
    add_text(s, "→", 8.25, 2.0, 0.6, 0.4, size=28, color=(14, 116, 144), bold=True)
    add_text(s, "Prediction probabilities", 9.0, 2.0, 3.4, 0.4, size=20, bold=True)
    add_text(s, "If a model says 90%, roughly 90% of similar cases should be correct. Calibration error measures the gap between reported probability and observed frequency.", 1.05, 3.25, 11.1, 1.0, size=23, color=(30, 41, 59))

    s = slide("2. Research question and hypothesis")
    add_text(s, "Question", 0.8, 1.6, 2.0, 0.4, size=22, color=(14, 116, 144), bold=True)
    add_text(s, "Does tracing a synthetic PTB-XL paper page change classwise calibration, and does the change differ by diagnosis?", 0.8, 2.1, 11.5, 0.9, size=26, bold=True)
    add_text(s, "Hypothesis", 0.8, 3.55, 2.0, 0.4, size=22, color=(14, 116, 144), bold=True)
    add_text(s, "The same digitization artifacts may affect probability reliability differently across diagnoses and across voltage representations.", 0.8, 4.05, 11.5, 0.8, size=23)

    s = slide("3. Data and model")
    add_text(s, "Dataset: PTB-XL 1.0.3", 0.8, 1.45, 5.5, 0.45, size=22, bold=True)
    add_text(s, "• 12-lead ECGs\n• 100 Hz, 10 seconds\n• 5 diagnostic superclasses\n• 17,084 train | 2,146 validation | 2,158 test", 0.95, 2.0, 5.4, 2.0, size=20)
    add_text(s, "Classifier", 7.0, 1.45, 3, 0.45, size=22, bold=True)
    add_text(s, "Compact 1D residual network\nBCE + AdamW\nSeed 42\nCheckpoint selected on fold 9", 7.1, 2.0, 5.2, 1.8, size=20)

    s = slide("4. The controlled comparison")
    add_text(s, "Same data • same seed • same architecture • same split", 1.0, 1.35, 11, 0.45, size=22, color=(14, 116, 144), bold=True)
    add_text(s, "Arm A\nPer-lead standardized\nRemoves absolute voltage scale", 1.0, 2.5, 4.2, 1.5, size=24, bold=True)
    add_text(s, "VS", 6.15, 2.85, 1.2, 0.5, size=26, color=(220, 38, 38), bold=True)
    add_text(s, "Arm B\nRaw millivolts\nKeeps absolute voltage scale", 8.0, 2.5, 4.2, 1.5, size=24, bold=True)
    add_text(s, "Primary question: does the representation change vulnerability to tracing?", 1.0, 5.25, 11, 0.5, size=22)

    s = slide("5. Paper-to-digital pipeline")
    add_text(s, "Clean waveform", 0.65, 1.5, 2.4, 0.4, size=20, bold=True)
    add_text(s, "→", 3.0, 1.5, 0.5, 0.4, size=28, color=(14, 116, 144), bold=True)
    add_text(s, "Synthetic paper page\nwrinkles + creases", 3.65, 1.35, 2.9, 0.8, size=20, bold=True)
    add_text(s, "→", 6.75, 1.5, 0.5, 0.4, size=28, color=(14, 116, 144), bold=True)
    add_text(s, "Trace recovery\ncolumn max-finder", 7.4, 1.35, 2.7, 0.8, size=20, bold=True)
    add_text(s, "→", 10.25, 1.5, 0.5, 0.4, size=28, color=(14, 116, 144), bold=True)
    add_text(s, "Paired evaluation", 10.8, 1.5, 2.1, 0.4, size=20, bold=True)
    add_text(s, "The same test record is scored twice: clean versus traced. Average trace SNR: 9.06 dB.", 1.0, 4.15, 11.3, 0.7, size=23)

    s = slide("6. Main result: calibration changes by class")
    s.shapes.add_picture(str(ASSETS / "ece_by_class.png"), Inches(0.65), Inches(1.25), width=Inches(12.0))
    add_text(s, "Positive = worse probability calibration after tracing", 0.8, 6.65, 11, 0.3, size=14, color=(71, 85, 105))

    s = slide("7. Main finding: raw millivolts are more vulnerable")
    s.shapes.add_picture(str(ASSETS / "contrast.png"), Inches(0.65), Inches(1.25), width=Inches(12.0))
    add_text(s, "Conduction disturbance contrast: +0.161 (patient-bootstrap interval ≈ 0.136–0.171)", 0.85, 6.62, 11.5, 0.35, size=17, color=(153, 27, 27), bold=True)

    s = slide("8. What does this mean?")
    add_text(s, "For conduction disturbance:", 0.85, 1.45, 5, 0.45, size=24, bold=True)
    add_text(s, "Raw-mV model:      +0.142 ΔECE\nStandardized model: −0.019 ΔECE\nDifference:          +0.161", 1.0, 2.2, 6.0, 1.8, size=25)
    add_text(s, "This is not just a reconstruction-quality story. CD records had similar SNR to non-CD records, and oracle amplitude matching removed only part of the shift.", 7.2, 2.0, 5.0, 2.0, size=21)

    s = slide("9. Did we improve the probability error?")
    s.shapes.add_picture(str(ASSETS / "correction.png"), Inches(0.65), Inches(1.25), width=Inches(12.0))
    add_text(s, "A validation-fitted recalibration map reduces the shift near zero without retraining the classifier.", 0.8, 6.62, 11.8, 0.35, size=17, color=(5, 100, 70), bold=True)

    s = slide("10. Novelty, contribution, and limits")
    add_text(s, "Contribution", 0.8, 1.3, 2.5, 0.4, size=22, color=(14, 116, 144), bold=True)
    add_text(s, "• Paired classwise calibration study after ECG paper digitization\n• Shows dependence on voltage representation\n• Separates calibration from AUROC and reconstruction quality\n• Demonstrates a practical recalibration fix", 0.95, 1.85, 6.0, 2.4, size=20)
    add_text(s, "Limits", 7.2, 1.3, 2.5, 0.4, size=22, color=(14, 116, 144), bold=True)
    add_text(s, "• Synthetic creased pages only\n• One primary seed and architecture\n• Real color scans not yet scored\n• Not a clinical deployment claim", 7.35, 1.85, 5.1, 2.2, size=20)
    add_text(s, "Take-home: paper digitization can preserve useful ranking while quietly damaging probability reliability.", 1.0, 5.4, 11.2, 0.65, size=25, color=(18, 59, 93), bold=True)

    prs.save(path)
    return path


if __name__ == "__main__":
    save_charts()
    pdf = make_detailed_pdf()
    ppt = make_ppt()
    print(pdf)
    print(ppt)
